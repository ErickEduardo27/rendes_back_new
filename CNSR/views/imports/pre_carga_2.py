from django.http import JsonResponse
import openpyxl
from datetime import datetime
from dateutil.relativedelta import relativedelta
from Asistencial.models import pacientes, modalidades, tipoPacientes, pacientesDialisis, etiologia, unidadesActuales,periodoIpress,usuarioIpress

def pre_carga_f1(request):
    if request.method == 'POST' and request.FILES['file']:

        #*************-----**************#
        id_modalidad = request.GET.get('id_modalidad')
        #id_tipo_paciente = request.GET.get('id_tipo_paciente')
        id_periodo_ipress = request.GET.get('id_periodo_ipress')
        id_usuario_ipress = request.GET.get('id_usuario_ipress')

        file = request.FILES['file']
        try:
            workbook = openpyxl.load_workbook(file, read_only=True, data_only=True)
            contador=0
            documentos = []
            indicador = 0
            # Iterar sobre las hojas del libro de trabajo comenzando desde la segunda hoja
            for sheet in workbook.worksheets[1:]:
                if contador == 0 :
                    for row in sheet.iter_rows(min_row=8, min_col=2,max_col=32, values_only=True):

                        tipo_documento, documento, autogenerado, nombre_paciente,edad,genero_id,genero,grado_instruccion_id,grado_instruccion, etiologia_general_id,etiologia_general,etiologia_especifica_id,etiologia_especifica , enf_ateroesclerotica_cardiaca, enf_insuficiencia_cardiaca_congestiva,enf_vascular_periferica,enf_cerebro_vascular,enf_cancer,enf_diabetes, enf_hipertension,enf_tuberculosis,enf_otra,modalidad_inicio_trr_id,modalidad_inicio_trr,fecha_inicio_trr,subsistema_salud_id,subsistema_salud,tipo_acceso_id,tipo_acceso,fecha_creacion_acceso,fecha_primer_ingreso  = row

                        hoy = datetime.today()
                        fecha_nacimiento = hoy - relativedelta(years=edad)
                        

                        if etiologia_especifica_id != "" and documento is not None : 
                            paciente_existente = pacientes.objects.filter(documento=documento).first()
                            documentos.append(documento)
                            etiologia_actual=etiologia.objects.get(codigo=etiologia_especifica_id)
                            paciente_actual=""
                            if paciente_existente and etiologia_especifica_id != "" and etiologia_actual:
                                paciente_existente.tipo_documento = tipo_documento
                                paciente_existente.autogenerado = autogenerado
                                paciente_existente.paciente = nombre_paciente
                                paciente_existente.fecha_nacimiento = fecha_nacimiento
                                paciente_existente.genero = genero
                                paciente_existente.grado_instruccion = grado_instruccion
                                paciente_existente.id_modalidad = modalidades.objects.get(id_modalidad=id_modalidad)
                                paciente_existente.id_tipo_paciente = tipoPacientes.objects.get(id_tipo_paciente=3)
                                paciente_existente.save()
                                paciente_actual=paciente_existente
                                paciente_dialisis_nuevo=pacientesDialisis(
                                    id_paciente = pacientes.objects.get(id_paciente=paciente_existente.id_paciente),
                                    id_etiologia = etiologia.objects.get(codigo='A.1.'),
                                    enf_ateroesclerotica_cardiaca=enf_ateroesclerotica_cardiaca,
                                    enf_insuficiencia_cardiaca_congestiva=enf_insuficiencia_cardiaca_congestiva,
                                    enf_vascular_periferica=enf_vascular_periferica,
                                    enf_cerebro_vascular=enf_cerebro_vascular,
                                    enf_cancer=enf_cancer,
                                    enf_diabetes=enf_diabetes,
                                    enf_hipertension=enf_hipertension,
                                    enf_tuberculosis=enf_tuberculosis, 
                                    enf_otra=enf_otra,
                                    id_periodo_ipress=periodoIpress.objects.get(id_periodo_ipress=id_periodo_ipress),
                                    id_usuario_ipress=usuarioIpress.objects.get(id_usuario_ipress=id_usuario_ipress),
                                    fecha_primer_ingreso=fecha_primer_ingreso,
                                    fecha_creacion_acceso=fecha_creacion_acceso,
                                    tipo_acceso=tipo_acceso,
                                    subsistema_salud=subsistema_salud,
                                    fecha_inicio_trr=fecha_inicio_trr,
                                    modalidad_inicio_trr=modalidad_inicio_trr
                                )
                                paciente_dialisis_nuevo.save()
                            else:
                                if etiologia_especifica_id != "" and etiologia_actual:
                                    paciente_nuevo = pacientes(
                                        documento=documento,
                                        tipo_documento=tipo_documento,
                                        autogenerado=autogenerado,
                                        paciente=nombre_paciente,
                                        fecha_nacimiento=fecha_nacimiento,
                                        genero=genero,
                                        grado_instruccion=grado_instruccion,
                                        id_modalidad=modalidades.objects.get(id_modalidad=id_modalidad),
                                        id_tipo_paciente=tipoPacientes.objects.get(id_tipo_paciente=1)
                                    )
                                    paciente_nuevo.save()
                                    paciente_actual=paciente_nuevo
                                    paciente_dialisis_nuevo=pacientesDialisis(
                                        id_paciente = pacientes.objects.get(id_paciente=paciente_nuevo.id_paciente),
                                        id_etiologia = etiologia.objects.get(codigo='A.1.'),
                                        enf_ateroesclerotica_cardiaca=enf_ateroesclerotica_cardiaca,
                                        enf_insuficiencia_cardiaca_congestiva=enf_insuficiencia_cardiaca_congestiva,
                                        enf_vascular_periferica=enf_vascular_periferica,
                                        enf_cerebro_vascular=enf_cerebro_vascular,
                                        enf_cancer=enf_cancer,
                                        enf_diabetes=enf_diabetes,
                                        enf_hipertension=enf_hipertension,
                                        enf_tuberculosis=enf_tuberculosis, 
                                        enf_otra=enf_otra,
                                        id_periodo_ipress=periodoIpress.objects.get(id_periodo_ipress=id_periodo_ipress),
                                        id_usuario_ipress=usuarioIpress.objects.get(id_usuario_ipress=id_usuario_ipress),
                                        fecha_primer_ingreso=fecha_primer_ingreso,
                                        fecha_creacion_acceso=fecha_creacion_acceso,
                                        tipo_acceso=tipo_acceso,
                                        subsistema_salud=subsistema_salud,
                                        fecha_inicio_trr=fecha_inicio_trr,
                                        modalidad_inicio_trr=modalidad_inicio_trr
                                    )
                                    paciente_dialisis_nuevo.save()
                    contador=contador+1
                if contador == 1:
                    indicador=0
                    for row in sheet.iter_rows(min_row=8, min_col=2, values_only=True):
                        nombres,fecha_ingreso,condicion_paciente,red,vacio_1,vacio_2,vhb_id,vhb,vhc_id,vhi_id,vhi,AcHBs_id,AcHBs,tipo_acceso_id,tipo_acceso,motivo_cambio_acceso_id,motivo_cambio_acceso,fecha_creacion_acceso,campos_vacio_1,campos_vacio_2,fecha_egreso1,campos_vacio_3,tipo_egreso1,fecha_reingreso1,fecha_egreso2,campos_vacio_4,tipo_egreso2,fecha_reingreso2,fecha_egreso3,campos_vacio_5,tipo_egreso3,fecha_reingreso3  = row
                        
                        unidad_actual=unidadesActuales(
                            id_paciente = pacientes.objects.get(documento=documentos[indicador]),
                            fecha_ingreso=fecha_ingreso,
                            condicion_paciente=condicion_paciente,
                            id_red = red.objects.get(red=red),
                            vhb=vhb,
                            vhc=vhc,
                            vhi=vhi,
                            AcHBs=AcHBs,
                            tipo_acceso=tipo_acceso,
                            motivo_cambio_acceso=motivo_cambio_acceso,
                            fecha_creacion_acceso=fecha_creacion_acceso
                        )
                        unidad_actual.save()
                        unidad_actual_id = unidad_actual.id_unidad_actual
                        if fecha_egreso1!="":
                            unidad_actual_detalles=unidadesActualesDetalles(
                                id_unidad_actual = unidadesActuales.objects.get(id_unidad_actual=unidad_actual_id),
                                fecha_egreso=fecha_egreso1,
                                tipo_egreso=tipo_egreso1,
                                fecha_reingreso = fecha_reingreso1,
                            )
                            unidad_actual_detalles.save()
                        if fecha_egreso2!="":
                                unidad_actual_detalles=unidadesActualesDetalles(
                                    id_unidad_actual = unidadesActuales.objects.get(id_unidad_actual=unidad_actual_id),
                                    fecha_egreso=fecha_egreso2,
                                    tipo_egreso=tipo_egreso2,
                                    fecha_reingreso = fecha_reingreso2,
                                )
                                unidad_actual_detalles.save()
                        if fecha_egreso3!="":
                                unidad_actual_detalles=unidadesActualesDetalles(
                                    id_unidad_actual = unidadesActuales.objects.get(id_unidad_actual=unidad_actual_id),
                                    fecha_egreso=fecha_egreso3,
                                    tipo_egreso=tipo_egreso3,
                                    fecha_reingreso = fecha_reingreso3,
                                )
                                unidad_actual_detalles.save()

                        indicador=indicador+1
                if contador == 2:
                    for row in sheet.iter_rows(min_row=8, min_col=2, values_only=True):
                        documento,nombres,tipo_acceso_vascular_id,tipo_acceso_vascular,fecha_evento,inicio_antmicrobial_id,inicio_antmicrobial,inicio_vancomicina_id,inicio_vancomicina,hemocultivo_positivo_id,hemocultivo_positivo,estado_acceso_vascular_id,estado_acceso_vascular,observaciones  = row
                            
                        paciente_existente = pacientes.objects.filter(documento=documento).first()

                        eventos_asociados=eventosAccesosVasculares(
                            id_paciente = pacientes.objects.get(id_paciente=paciente_existente.id),
                            id_usuario_ipress=usuarioIpress.objects.get(id_usuario_ipress=id_usuario_ipress),
                            tipo_acceso_vascular=tipo_acceso_vascular,
                            fecha_evento = fecha_evento,
                            inicio_antmicrobial=inicio_antmicrobial,
                            inicio_vancomicina=inicio_vancomicina,
                            hemocultivo_positivo=hemocultivo_positivo,
                            estado_acceso_vascular=estado_acceso_vascular,
                            observaciones=observaciones,
                            id_periodo_ipress=periodoIpress.objects.get(id_periodo_ipress=id_periodo_ipress)
                        )
                        eventos_asociados.save()
                if contador == 3:
                    for row in sheet.iter_rows(min_row=8, min_col=2, values_only=True):
                        documento,nombres,diagnostico,codigo_diagnostico,fecha_hospitalizacion,fecha_alta,fuente  = row
                            
                        paciente_existente = pacientes.objects.filter(documento=documento).first()
                        if paciente_existente:
                            morbilidad_hospitalaria=morbilidadesHospitalarias(
                                id_paciente = pacientes.objects.get(id_paciente=paciente_existente.id),
                                id_usuario_ipress=usuarioIpress.objects.get(id_usuario_ipress=id_usuario_ipress),
                                diagnostico=diagnostico,
                                codigo_diagnostico = codigo_diagnostico,
                                fecha_hospitalizacion=fecha_hospitalizacion,
                                fecha_alta=fecha_alta,
                                fuente=fuente,
                                id_periodo_ipress=periodoIpress.objects.get(id_periodo_ipress=id_periodo_ipress)
                            )
                            morbilidad_hospitalaria.save()
                
                if contador == 4:
                    indicador=0
                    for row in sheet.iter_rows(min_row=8, min_col=2, values_only=True):
                        nombres,Hb,calcio,fosforo,PTHi,Alb,calcio_corregido,ktv,tiempo_dialisis_id,tiempo_dialisis,eritoproyetina_id,eritoproyetina,hierro_id,hiperparatioidismo_id,hiperparatioidismo  = row
                            
                        paciente_existente = pacientes.objects.filter(documento=documentos[indicador]).first()
                        if paciente_existente:
                            resultados_clinicos=resultadosClinicos(
                                id_paciente = pacientes.objects.get(id_paciente=paciente_existente.id),
                                id_usuario_ipress=usuarioIpress.objects.get(id_usuario_ipress=id_usuario_ipress),
                                Hb=Hb,
                                calcio = calcio,
                                fosforo= fosforo,
                                PTHi=PTHi,
                                Alb=Alb,
                                calcio_corregido=calcio_corregido,
                                ktv=ktv,
                                tiempo_dialisis=tiempo_dialisis,
                                eritoproyetina=eritoproyetina,
                                hierro=hierro,
                                hiperparatioidismo=hiperparatioidismo,
                                id_periodo_ipress=periodoIpress.objects.get(id_periodo_ipress=id_periodo_ipress)
                            )
                            resultados_clinicos.save()
                        indicador=indicador+1

                if contador == 5:
                    for row in sheet.iter_rows(min_row=8, min_col=1, values_only=True):
                        control,periodo,salida_osmosis_ufc,anillo_circulacion_ufc,salida_osmosis_ue,anillo_circulacion_ue,maquina_1_ufc,maquina_2_ufc,maquina_1_ue,maquina_2_ue  = row
                        calidad_microbiologica=calidadMicrobiologicas(
                            id_usuario_ipress=usuarioIpress.objects.get(id_usuario_ipress=id_usuario_ipress),
                            control=control,
                            salida_osmosis_ufc = salida_osmosis_ufc,
                            anillo_circulacion_ufc= anillo_circulacion_ufc,
                            salida_osmosis_ue=salida_osmosis_ue,
                            anillo_circulacion_ue=anillo_circulacion_ue,
                            maquina_1_ufc=maquina_1_ufc,
                            maquina_2_ufc=maquina_2_ufc,
                            maquina_1_ue=maquina_1_ue,
                            maquina_2_ue=maquina_2_ue,
                            hierro=hierro,
                            id_periodo_ipress=periodoIpress.objects.get(id_periodo_ipress=id_periodo_ipress)
                        )
                        calidad_microbiologica.save()
                if contador == 6:
                    indicador=0
                    for row in sheet.iter_rows(min_row=8, min_col=2, values_only=True):
                        nombres,red,turno,frecuencia,nro_dosis1,fecha_dosis1,motivo_no_vacunacion1 ,nro_dosis2,fecha_dosis2,motivo_no_vacunacion2,nro_dosis3,fecha_dosis3,motivo_no_vacunacion3,nro_dosis4,fecha_dosis4,motivo_no_vacunacion4 = row
                        paciente_existente = pacientes.objects.filter(documento=documentos[indicador]).first()
                        if paciente_existente:
                            vacunacion=vacunaciones(
                                id_paciente = pacientes.objects.get(id_paciente=paciente_existente.id),
                                id_usuario_ipress=usuarioIpress.objects.get(id_usuario_ipress=id_usuario_ipress),
                                id_red= red.objects.get(red=red),
                                turno = turno,
                                frecuencia= frecuencia,
                                id_periodo_ipress=periodoIpress.objects.get(id_periodo_ipress=id_periodo_ipress)
                            )
                            vacunacion.save()
                            vacunacion_id=calidad_microbiologica.id
                            if nro_dosis1!="":
                                vacunacion_detalle=vacunacionesDetalles(
                                    id_vacunacion=vacunaciones.objects.get(id_vacunacion=vacunacion_id),
                                    tipo_vacuna="Vacunación Contra Hepatitis B",
                                    nro_dosis=nro_dosis1,
                                    fecha_dosis=fecha_dosis1,
                                    motivo_no_vacunacion=motivo_no_vacunacion1
                                )
                            if nro_dosis2!="":
                                vacunacion_detalle=vacunacionesDetalles(
                                    id_vacunacion=vacunaciones.objects.get(id_vacunacion=vacunacion_id),
                                    tipo_vacuna="Vacunación contra COVID-19",
                                    nro_dosis=nro_dosis2,
                                    fecha_dosis=fecha_dosis2,
                                    motivo_no_vacunacion=motivo_no_vacunacion2
                                )
                            if nro_dosis3!="":
                                vacunacion_detalle=vacunacionesDetalles(
                                    id_vacunacion=vacunaciones.objects.get(id_vacunacion=vacunacion_id),
                                    tipo_vacuna="Vacunación contra Influenza",
                                    nro_dosis=nro_dosis3,
                                    fecha_dosis=fecha_dosis3,
                                    motivo_no_vacunacion=motivo_no_vacunacion3
                                )
                            if nro_dosis4!="":
                                vacunacion_detalle=vacunacionesDetalles(
                                    id_vacunacion=vacunaciones.objects.get(id_vacunacion=vacunacion_id),
                                    tipo_vacuna="Vacunación contra Neumococo",
                                    nro_dosis=nro_dosis4,
                                    fecha_dosis=fecha_dosis4,
                                    motivo_no_vacunacion=motivo_no_vacunacion4
                                )

                        indicador=indicador+1
                contador=contador+1
            return JsonResponse({'message': 'Datos guardados exitosamente.'})
        except Exception as e:
            return JsonResponse({'error': str(e)}, status=500)
    else:
        return JsonResponse({'error': 'Debes enviar un archivo.'}, status=400)