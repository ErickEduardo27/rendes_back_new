from django.http import JsonResponse
import openpyxl
from datetime import datetime
from dateutil.relativedelta import relativedelta
from Asistencial.models import pacientes, modalidades, tipoPacientes, pacientesDialisis, etiologia, unidadesActuales,periodoIpress,usuarioIpress,red,unidadesActualesDetalles

def pre_carga_f1(request):
    if request.method == 'POST' and request.FILES['file']:

        #*************-----**************#
        id_modalidad = request.GET.get('id_modalidad')
        #id_tipo_paciente = request.GET.get('id_tipo_paciente')
        id_periodo_ipress = request.GET.get('id_periodo_ipress')
        id_usuario_ipress = request.GET.get('id_usuario_ipress')

        file = request.FILES['file']
        try:
            workbook = openpyxl.load_workbook(file, read_only=True, data_only=True)
            contador=0
            documentos = []
            indicador = 0
            # Iterar sobre las hojas del libro de trabajo comenzando desde la segunda hoja
            for sheet in workbook.worksheets[1:]:
                if contador == 0 :
                    for row in sheet.iter_rows(min_row=8, min_col=2,max_col=32, values_only=True):

                        tipo_documento, documento, autogenerado, nombre_paciente,edad,genero_id,genero,grado_instruccion_id,grado_instruccion, etiologia_general_id,etiologia_general,etiologia_especifica_id,etiologia_especifica , enf_ateroesclerotica_cardiaca, enf_insuficiencia_cardiaca_congestiva,enf_vascular_periferica,enf_cerebro_vascular,enf_cancer,enf_diabetes, enf_hipertension,enf_tuberculosis,enf_otra,modalidad_inicio_trr_id,modalidad_inicio_trr,fecha_inicio_trr,subsistema_salud_id,subsistema_salud,tipo_acceso_id,tipo_acceso,fecha_creacion_acceso,fecha_primer_ingreso  = row

                        hoy = datetime.today()

                        if etiologia_especifica_id != "" and documento is not None :
                            fecha_nacimiento = hoy - relativedelta(years=edad)
                            paciente_existente = pacientes.objects.filter(documento=documento).first()
                            documentos.append(documento)
                            etiologia_actual=etiologia.objects.get(codigo=etiologia_especifica_id)
                            if paciente_existente and etiologia_especifica_id != "" and etiologia_actual:
                                paciente_existente.tipo_documento = tipo_documento
                                paciente_existente.autogenerado = autogenerado
                                paciente_existente.paciente = nombre_paciente
                                paciente_existente.fecha_nacimiento = fecha_nacimiento
                                paciente_existente.genero = genero
                                paciente_existente.grado_instruccion = grado_instruccion
                                paciente_existente.id_modalidad = modalidades.objects.get(id_modalidad=id_modalidad)
                                paciente_existente.save()
                                paciente_actual=paciente_existente
                            else:
                                if etiologia_especifica_id != "" and etiologia_actual and documento is not None:
                                    paciente_nuevo = pacientes(
                                        documento=documento,
                                        tipo_documento=tipo_documento,
                                        autogenerado=autogenerado,
                                        paciente=nombre_paciente,
                                        fecha_nacimiento=fecha_nacimiento,
                                        genero=genero,
                                        grado_instruccion=grado_instruccion,
                                        id_modalidad=modalidades.objects.get(id_modalidad=id_modalidad),
                                    )
                                    paciente_nuevo.save()
                                    paciente_actual=paciente_dialisis_nuevo
                            paciente_dialisis = pacientesDialisis.objects.filter(
                                id_paciente=paciente_actual.id_paciente,
                                id_periodo_ipress=id_periodo_ipress,
                                id_usuario_ipress=id_usuario_ipress
                            ).first()

                            if paciente_dialisis and etiologia_especifica_id != "" and etiologia_actual:
                                paciente_dialisis.id_paciente = pacientes.objects.get(id_paciente=paciente_actual.id_paciente)
                                paciente_dialisis.id_etiologia = etiologia.objects.get(codigo=etiologia_especifica_id)
                                paciente_dialisis.enf_ateroesclerotica_cardiaca = enf_ateroesclerotica_cardiaca
                                paciente_dialisis.enf_insuficiencia_cardiaca_congestiva = enf_insuficiencia_cardiaca_congestiva
                                paciente_dialisis.enf_vascular_periferica = enf_vascular_periferica
                                paciente_dialisis.enf_cerebro_vascular = enf_cerebro_vascular
                                paciente_dialisis.enf_cancer = enf_cancer
                                paciente_dialisis.enf_diabetes = enf_diabetes
                                paciente_dialisis.enf_hipertension = enf_hipertension
                                paciente_dialisis.enf_tuberculosis = enf_tuberculosis
                                paciente_dialisis.enf_otra = enf_otra
                                paciente_dialisis.id_periodo_ipress=periodoIpress.objects.get(id_periodo_ipress=id_periodo_ipress)
                                paciente_dialisis.id_usuario_ipress=usuarioIpress.objects.get(id_usuario_ipress=id_usuario_ipress)
                                paciente_dialisis.fecha_primer_ingreso=fecha_primer_ingreso
                                paciente_dialisis.fecha_creacion_acceso=fecha_creacion_acceso
                                paciente_dialisis.tipo_acceso=tipo_acceso
                                paciente_dialisis.subsistema_salud=subsistema_salud
                                paciente_dialisis.fecha_inicio_trr=fecha_inicio_trr
                                paciente_dialisis.modalidad_inicio_trr=modalidad_inicio_trr
                                paciente_dialisis.save()

                            if paciente_dialisis is None and etiologia_especifica_id != "" and etiologia_actual:
                                paciente_dialisis_nuevo=pacientesDialisis(
                                        id_paciente = pacientes.objects.get(id_paciente=paciente_actual.id_paciente),
                                        id_etiologia = etiologia.objects.get(codigo=etiologia_especifica_id),
                                        enf_ateroesclerotica_cardiaca=enf_ateroesclerotica_cardiaca,
                                        enf_insuficiencia_cardiaca_congestiva=enf_insuficiencia_cardiaca_congestiva,
                                        enf_vascular_periferica=enf_vascular_periferica,
                                        enf_cerebro_vascular=enf_cerebro_vascular,
                                        enf_cancer=enf_cancer,
                                        enf_diabetes=enf_diabetes,
                                        enf_hipertension=enf_hipertension,
                                        enf_tuberculosis=enf_tuberculosis, 
                                        enf_otra=enf_otra,
                                        id_periodo_ipress=periodoIpress.objects.get(id_periodo_ipress=id_periodo_ipress),
                                        id_usuario_ipress=usuarioIpress.objects.get(id_usuario_ipress=id_usuario_ipress),
                                        fecha_primer_ingreso=fecha_primer_ingreso,
                                        fecha_creacion_acceso=fecha_creacion_acceso,
                                        tipo_acceso=tipo_acceso,
                                        subsistema_salud=subsistema_salud,
                                        fecha_inicio_trr=fecha_inicio_trr,
                                        modalidad_inicio_trr=modalidad_inicio_trr
                                )
                                paciente_dialisis_nuevo.save()

                if contador == 1:
                    indicador=0
                    for row in sheet.iter_rows(min_row=8, min_col=2,max_col=35, values_only=True):
                       
                        nombres,fecha_ingreso,condicion_paciente_1,condicion_paciente,redes,libre1,libre2,vhb_id,vhb,vhc_id,vhc,vhi_id,vhi,libre3,libre4,AcHBs_id,AcHBs,tipo_acceso_id,tipo_acceso,motivo_cambio_acceso_id,motivo_cambio_acceso,fecha_creacion_acceso,fecha_egreso1,campos_vacio_3,tipo_egreso1,fecha_reingreso1,fecha_egreso2,campos_vacio_4,tipo_egreso2,fecha_reingreso2,fecha_egreso3,campos_vacio_5,tipo_egreso3,fecha_reingreso3  = row
                        print(len(documentos))
                        print(indicador)
                        if len(documentos)>0 and indicador<len(documentos):
                            
                            paciente_existente = pacientes.objects.filter(documento=documentos[indicador]).first()
                            if paciente_existente:
                                unidad_actual = unidadesActuales.objects.filter(
                                        id_paciente=paciente_existente.id_paciente,
                                        id_periodo_ipress=id_periodo_ipress,
                                        id_usuario_ipress=id_usuario_ipress
                                    ).first()
                                if unidad_actual:
                                    unidad_actual.id_paciente = pacientes.objects.get(id_paciente=paciente_existente.id_paciente)
                                    unidad_actual.fecha_ingreso=fecha_ingreso
                                    unidad_actual.id_red = red.objects.get(id_red=1)
                                    unidad_actual.VHB=vhb
                                    unidad_actual.VHC=vhc
                                    unidad_actual.VHI=vhi
                                    unidad_actual.AcHBs=AcHBs
                                    unidad_actual.tipo_acceso=tipo_acceso
                                    unidad_actual.motivo_cambio_acceso=motivo_cambio_acceso
                                    unidad_actual.id_usuario_ipress=usuarioIpress.objects.get(id_usuario_ipress=id_usuario_ipress)
                                    unidad_actual.id_periodo_ipress=periodoIpress.objects.get(id_periodo_ipress=id_periodo_ipress)
                                    unidad_actual.id_tipo_paciente=tipoPacientes.objects.get(id_tipo_paciente=condicion_paciente_1)
                                    unidad_actual.fecha_creacion_acceso=fecha_creacion_acceso
                                    unidad_actual.save()

                                    unidad_actual_id = unidad_actual.id_unidad_actual
                                    unidad_actual_detalles = unidadesActualesDetalles.objects.filter(id_unidad_actual=unidad_actual_id)
                                    unidad_actual_detalle_actual=unidad_actual_detalles
                                    
                                else:
                                    unidad_actual=unidadesActuales(
                                        id_paciente = pacientes.objects.get(id_paciente=paciente_existente.id_paciente),
                                        fecha_ingreso=fecha_ingreso,
                                        id_red = red.objects.get(id_red=1),
                                        VHB=vhb,
                                        VHC=vhc,
                                        VHI=vhi,
                                        AcHBs=AcHBs,
                                        tipo_acceso=tipo_acceso,
                                        motivo_cambio_acceso=motivo_cambio_acceso,
                                        fecha_creacion_acceso=fecha_creacion_acceso,
                                        id_usuario_ipress=usuarioIpress.objects.get(id_usuario_ipress=id_usuario_ipress),
                                        id_periodo_ipress=periodoIpress.objects.get(id_periodo_ipress=id_periodo_ipress),
                                        id_tipo_paciente=tipoPacientes.objects.get(id_tipo_paciente=1)
                                    )
                                    unidad_actual.save()
                                    unidad_actual_id = unidad_actual.id_unidad_actual
                                    unidad_actual_detalles = unidadesActualesDetalles.objects.filter(id_unidad_actual=unidad_actual_id)
                                    unidad_actual_detalle_actual=unidad_actual_detalles
                                """ print(unidad_actual_detalle_actual[0]) """
                                if len(unidad_actual_detalle_actual)>0 and fecha_egreso1!="":
                                        unidad_actual_detalles = unidadesActualesDetalles.objects.filter(id_unidad_actual_detalle=unidad_actual_detalle_actual[0].id_unidad_actual_detalle).first()
                                        unidad_actual_detalles.id_unidad_actual = unidadesActuales.objects.get(id_unidad_actual=unidad_actual_id)
                                        unidad_actual_detalles.fecha_egreso=fecha_egreso1
                                        unidad_actual_detalles.tipo_egreso=tipo_egreso1
                                        unidad_actual_detalles.fecha_reingreso = fecha_reingreso1
                                        unidad_actual_detalles.save()
                                """ print("error") """
                                if len(unidad_actual_detalle_actual)>0 is None and fecha_egreso1!="":
                                        unidad_actual_detalles=unidadesActualesDetalles(
                                            id_unidad_actual = unidadesActuales.objects.get(id_unidad_actual=unidad_actual_id),
                                            fecha_egreso=fecha_egreso1,
                                            tipo_egreso=tipo_egreso1,
                                            fecha_reingreso = fecha_reingreso1,
                                        )
                                        unidad_actual_detalles.save()

                                if len(unidad_actual_detalle_actual)>1 and fecha_egreso2!="":
                                        unidad_actual_detalles = unidadesActualesDetalles.objects.filter(id_unidad_actual_detalle=unidad_actual_detalle_actual[1].id_unidad_actual_detalle).first()
                                        unidad_actual_detalles.id_unidad_actual = unidadesActuales.objects.get(id_unidad_actual=unidad_actual_id)
                                        unidad_actual_detalles.fecha_egreso=fecha_egreso2
                                        unidad_actual_detalles.tipo_egreso=tipo_egreso2
                                        unidad_actual_detalles.fecha_reingreso = fecha_reingreso2
                                        unidad_actual_detalles.save()
                                if len(unidad_actual_detalle_actual)>1 is None and fecha_egreso2!="":
                                        unidad_actual_detalles=unidadesActualesDetalles(
                                            id_unidad_actual = unidadesActuales.objects.get(id_unidad_actual=unidad_actual_id),
                                            fecha_egreso=fecha_egreso2,
                                            tipo_egreso=tipo_egreso2,
                                            fecha_reingreso = fecha_reingreso2,
                                        )
                                        unidad_actual_detalles.save()

                                if len(unidad_actual_detalle_actual)>2 and fecha_egreso3!="":
                                        unidad_actual_detalles = unidadesActualesDetalles.objects.filter(id_unidad_actual_detalle=unidad_actual_detalle_actual[2].id_unidad_actual_detalle).first()
                                        unidad_actual_detalles.id_unidad_actual = unidadesActuales.objects.get(id_unidad_actual=unidad_actual_id)
                                        unidad_actual_detalles.fecha_egreso=fecha_egreso3
                                        unidad_actual_detalles.tipo_egreso=tipo_egreso3
                                        unidad_actual_detalles.fecha_reingreso = fecha_reingreso3
                                        unidad_actual_detalles.save()
                                if len(unidad_actual_detalle_actual)>2 is None and fecha_egreso2!="":
                                        unidad_actual_detalles=unidadesActualesDetalles(
                                            id_unidad_actual = unidadesActuales.objects.get(id_unidad_actual=unidad_actual_id),
                                            fecha_egreso=fecha_egreso3,
                                            tipo_egreso=tipo_egreso3,
                                            fecha_reingreso = fecha_reingreso3,
                                        )
                                        unidad_actual_detalles.save()
                                if(indicador<len(documentos)):
                                    indicador=indicador+1
                contador=contador+1
            return JsonResponse({'message': 'Datos guardados exitosamente.'})
        except Exception as e:
            return JsonResponse({'error': str(e)}, status=500)
    else:
        return JsonResponse({'error': 'Debes enviar un archivo.'}, status=400)